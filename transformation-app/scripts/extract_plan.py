#!/usr/bin/env python3
"""
Regenerates js/data/planData.js from the source Excel workbook.

Usage:
    python3 scripts/extract_plan.py /path/to/1700_Calorie_Indian_Cutting_Plan.xlsx

Run this after editing the Excel (add/change meals, workout days, grocery
items, or targets) to refresh the app's plan data. It never touches any
logged tracking data — only the read-only PLAN object the app treats as the
source of truth.

Requires: pip install openpyxl
"""
import json
import re
import sys
from pathlib import Path

DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
MEAL_ORDER = ["Breakfast", "Mid-morning", "Fruit", "Lunch", "Snack", "Dinner", "Treat"]


def extract(path):
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)

    # ---- 7-Day Meal Plan ----
    ws = wb["7-Day Meal Plan"]
    meals_raw = []
    mid = 1
    for r in ws.iter_rows(min_row=2, values_only=True):
        day = r[0]
        if day is None or day == "Note":
            continue
        meals_raw.append(
            {
                "id": f"m{mid}",
                "day": day,
                "meal": r[1],
                "food": r[2],
                "quantity": r[3],
                "calories": r[4],
                "protein": r[5],
                "carbs": r[6],
                "fat": r[7],
                "fiber": r[8],
                "notes": r[9] or "",
            }
        )
        mid += 1

    meal_plan = {day: [] for day in DAYS}
    for m in meals_raw:
        meal_plan.setdefault(m["day"], []).append(m)
    for day in DAYS:
        meal_plan[day].sort(key=lambda m: MEAL_ORDER.index(m["meal"]) if m["meal"] in MEAL_ORDER else 99)

    # ---- Targets & Guidelines ----
    ws = wb["Targets & Guidelines"]
    targets_raw = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        targets_raw[r[0]] = r[1]

    def numify(n):
        return int(n) if float(n).is_integer() else round(float(n), 2)

    def rng(s):
        nums = re.findall(r"[\d.]+", str(s).replace("–", "-").replace("—", "-"))
        return [numify(n) for n in nums]

    def target_entry(key, label, default_min, default_max, default_val, unit):
        guidance = targets_raw.get(key, "")
        nums = rng(guidance)
        lo = nums[0] if nums else default_min
        hi = nums[1] if len(nums) > 1 else (nums[0] if nums else default_max)
        default = default_val if default_val is not None else numify(round((lo + hi) / 2, 1))
        return {"min": lo, "max": hi, "default": default, "unit": unit, "label": label, "guidance": guidance}

    daily_targets = {
        "calories": target_entry("Calories", "Calories", 1700, 1700, 1700, "kcal"),
        "protein": target_entry("Protein", "Protein", 150, 170, 160, "g"),
        "fiber": target_entry("Fiber", "Fiber", 25, 35, 30, "g"),
        "water": target_entry("Water", "Water", 2.5, 3.5, 3.0, "L"),
        "vegetables": target_entry("Vegetables", "Vegetables", 300, 500, 400, "g"),
        "darkChocolate": target_entry("Dark chocolate", "Dark chocolate", 10, 15, 12, "g"),
        "whey": target_entry("Whey", "Whey", 1, 1, 1, "scoop"),
    }
    handled = {"Calories", "Protein", "Fiber", "Water", "Vegetables", "Dark chocolate", "Whey"}
    guidelines = {k: v for k, v in targets_raw.items() if k not in handled}

    # ---- Workout & Cardio ----
    ws = wb["Workout & Cardio"]
    workout_plan = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        text = r[1] or ""
        if ":" in text:
            label, rest = text.split(":", 1)
            exercises = [e.strip() for e in rest.split(",") if e.strip()]
        else:
            label, exercises = text.strip(), []
        workout_plan[r[0]] = {
            "day": r[0],
            "label": label.strip(),
            "exercises": exercises,
            "cardio": r[2],
            "duration": r[3],
            "isRestDay": r[0] == "Sunday",
            "isOptional": r[0] == "Saturday",
        }

    # ---- Weekly Grocery List ----
    ws = wb["Weekly Grocery List"]
    grocery = []
    gid = 1
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        grocery.append({"id": f"g{gid}", "category": r[0], "item": r[1], "amount": r[2], "notes": r[3] or ""})
        gid += 1

    grocery_by_category = {}
    for g in grocery:
        grocery_by_category.setdefault(g["category"], []).append(g)

    return {
        "days": DAYS,
        "mealOrder": MEAL_ORDER,
        "mealPlan": meal_plan,
        "workoutPlan": workout_plan,
        "dailyTargets": daily_targets,
        "guidelines": guidelines,
        "grocery": grocery,
        "groceryByCategory": grocery_by_category,
    }


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    xlsx_path = sys.argv[1]
    data = extract(xlsx_path)

    repo_root = Path(__file__).resolve().parent.parent
    out_path = repo_root / "js" / "data" / "planData.js"
    js = (
        "// AUTO-GENERATED by scripts/extract_plan.py — do not hand-edit generated blocks below\n"
        "// without re-running the script, or your changes will be overwritten next time.\n"
        "// Safe to hand-edit in place if you'd rather skip Excel entirely.\n\n"
        "export const PLAN = " + json.dumps(data, indent=2, ensure_ascii=False) + ";\n"
    )
    out_path.write_text(js, encoding="utf-8")
    print(f"Wrote {out_path} ({len(js)} bytes)")
    print(f"  Meals: {sum(len(v) for v in data['mealPlan'].values())}")
    print(f"  Workout days: {len(data['workoutPlan'])}")
    print(f"  Grocery items: {len(data['grocery'])}")


if __name__ == "__main__":
    main()
